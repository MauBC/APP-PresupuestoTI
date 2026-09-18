from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
)
from openpyxl.utils import (
    get_column_letter,
)

from app.config.budget_module_config import (
    BudgetModule,
    BudgetModuleConfig,
)
from app.config.capex_schema import (
    CAPEX_EXCEL_SHEET,
    CAPEX_INTERNAL_TO_RAW,
)


HEADER_FILL = "2F7650"
HEADER_FONT = "FFFFFF"


class BudgetExcelExportError(
    RuntimeError
):
    pass


@dataclass(
    frozen=True
)
class BudgetExcelExportResult:
    path: str
    module: str
    sheet_name: str
    row_count: int
    column_count: int


class BudgetExcelExportService:
    def __init__(
        self,
        repository,
        *,
        module_config:
            BudgetModuleConfig | None
            = None,
    ):
        self._repository = repository

        self._config = (
            module_config
            if module_config is not None
            else repository.module_config
        )

    @property
    def module_config(
        self,
    ) -> BudgetModuleConfig:
        return self._config

    @property
    def export_columns(
        self,
    ) -> tuple[str, ...]:
        return tuple(
            self._config
            .insert_columns
        )

    def _sheet_name(
        self,
    ) -> str:
        if (
            self._config.module
            == BudgetModule.CAPEX
        ):
            return CAPEX_EXCEL_SHEET

        return "OPEX 2027"

    def _header_for_column(
        self,
        column: str,
    ) -> str:
        if (
            self._config.module
            == BudgetModule.CAPEX
        ):
            return (
                CAPEX_INTERNAL_TO_RAW
                .get(
                    column,
                    column,
                )
            )

        return column

    @staticmethod
    def _excel_value(
        value,
    ):
        if isinstance(
            value,
            Decimal,
        ):
            return float(
                value
            )

        return value

    @staticmethod
    def _normalize_target(
        destination,
    ) -> Path:
        target = Path(
            destination
        ).expanduser()

        if (
            target.suffix.lower()
            != ".xlsx"
        ):
            target = (
                target.with_suffix(
                    ".xlsx"
                )
            )

        return target

    def export(
        self,
        destination,
    ) -> BudgetExcelExportResult:
        target = (
            self._normalize_target(
                destination
            )
        )

        try:
            source_rows = tuple(
                self._repository
                .get_all_rows()
            )

            rows = tuple(
                row
                for row in source_rows
                if (
                    row.get(
                        "habilitado",
                        True,
                    )
                    is not False
                )
            )

        except Exception as exc:
            raise BudgetExcelExportError(
                "No se pudieron leer "
                "los datos confirmados "
                "de BigQuery."
            ) from exc

        columns = (
            self.export_columns
        )

        if not columns:
            raise BudgetExcelExportError(
                "El modulo no tiene "
                "columnas exportables."
            )

        target.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        workbook = Workbook()

        worksheet = workbook.active

        worksheet.title = (
            self._sheet_name()
        )

        headers = tuple(
            self._header_for_column(
                column
            )
            for column in columns
        )

        worksheet.append(
            headers
        )

        for row in rows:
            worksheet.append(
                tuple(
                    self._excel_value(
                        row.get(
                            column
                        )
                    )
                    for column
                    in columns
                )
            )

        self._format_sheet(
            worksheet,
            columns=columns,
        )

        try:
            workbook.save(
                target
            )

        except Exception as exc:
            raise BudgetExcelExportError(
                "No se pudo guardar "
                "el archivo Excel."
            ) from exc

        finally:
            workbook.close()

        return (
            BudgetExcelExportResult(
                path=str(
                    target.resolve()
                ),
                module=(
                    self._config
                    .module
                    .value
                ),
                sheet_name=(
                    self._sheet_name()
                ),
                row_count=len(
                    rows
                ),
                column_count=len(
                    columns
                ),
            )
        )

    def _format_sheet(
        self,
        worksheet,
        *,
        columns,
    ):
        worksheet.freeze_panes = (
            "A2"
        )

        for cell in worksheet[1]:
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor=HEADER_FILL,
            )

            cell.font = Font(
                bold=True,
                color=HEADER_FONT,
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        worksheet.row_dimensions[
            1
        ].height = 32

        type_map = (
            self._config
            .insert_type_map
        )

        for (
            index,
            column,
        ) in enumerate(
            columns,
            start=1,
        ):
            value_type = (
                type_map.get(
                    column,
                    "STRING",
                )
            )

            header = (
                self._header_for_column(
                    column
                )
            )

            if value_type == "STRING":
                width = max(
                    14,
                    min(
                        32,
                        len(
                            str(
                                header
                            )
                        ) + 4,
                    ),
                )

            else:
                width = max(
                    14,
                    min(
                        20,
                        len(
                            str(
                                header
                            )
                        ) + 4,
                    ),
                )

            worksheet.column_dimensions[
                get_column_letter(
                    index
                )
            ].width = width

        if (
            worksheet.max_row
            >= 1
            and
            worksheet.max_column
            >= 1
        ):
            last_column = (
                get_column_letter(
                    worksheet.max_column
                )
            )

            worksheet.auto_filter.ref = (
                f"A1:"
                f"{last_column}"
                f"{worksheet.max_row}"
            )
