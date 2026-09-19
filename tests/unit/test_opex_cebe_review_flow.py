from dataclasses import replace
from decimal import Decimal
import json
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from app.config.budget_modules import OPEX_MODULE_CONFIG
from app.config.presupuesto_schema import MONTHS
from app.config.opex_smart_precision import OPEX_SMART_MONEY_QUANTUM
from app.services.presupuesto_workspace import PresupuestoWorkspace
from app.services.budget_excel_export_service import BudgetExcelExportService
from database.persistence.batch_builder import build_persistence_batch
from database.persistence.staging_builder import build_staging_rows

from app.models.opex_fx import OpexFxTable
from app.models.opex_master_data import (
    OpexAccountMasterRecord, OpexCebeMasterRecord, OpexMasterConflict,
    OpexMasterDataSnapshot, OpexMasterSource, OpexRecoverableMasterRecord,
)
from app.models.opex_smart_import import OpexSmartImportSheetOverride
from app.models.opex_smart_template import (
    OpexSmartTemplateWorkbook, OpexTemplateBudget, OpexTemplateDistribution,
)
from app.services import opex_smart_import_preparation_service as preparation


pytestmark = pytest.mark.unit


@pytest.mark.parametrize(("tipo", "monto"), [
    ("MENSUAL", "100"),
    ("MENSUAL", "0.001153901"),
    ("ANUAL", "0.001153901"),
    ("ANUAL", "0.000000002"),
    ("MENSUAL", "5954.3499999999985"),
])
def test_review_to_workspace_staging_and_export_preserves_amounts(tmp_path, monkeypatch, tipo, monto):
    source = tmp_path / "plantilla.xlsx"
    source.touch()
    account = OpexAccountMasterRecord("600000001", "SERVICIOS", "SOFTWARE", "GASTOS TI")
    first = OpexCebeMasterRecord(
        "51IC000000", "TESORERIA", "CORPORATIVO", "TESORERIA", "LIMA", "SEDE A", "SEGMENTO",
    )
    second = replace(first, sede_cg="SEDE B")
    metadata = OpexMasterSource("test.xlsx", "Sheet1", 1, 2, 2, 0)
    snapshot = OpexMasterDataSnapshot(
        accounts={account.numero_cuenta: account}, cebes={},
        recoverables={"51": OpexRecoverableMasterRecord("51", "2501", "EMPRESA", "PE")},
        account_conflicts={}, recoverable_conflicts={},
        cebe_conflicts={first.centro_beneficio: OpexMasterConflict(
            first.centro_beneficio, ("Sheet1:2", "Sheet1:3"), (first, second),
        )},
        account_source=metadata, cebe_source=metadata, recoverable_source=metadata,
    )
    budget = OpexTemplateBudget(
        "Licencias", "Software", "Proveedor", "USD", account.numero_cuenta,
        tipo, Decimal(monto),
        (OpexTemplateDistribution(8, "51IC000002", None, Decimal("1")),),
    )
    workbook = OpexSmartTemplateWorkbook(str(source), (budget,))
    monkeypatch.setattr(preparation.OpexSmartTemplateLoader, "load", lambda self, path: workbook)
    monkeypatch.setattr(preparation.OpexMasterDataLoader, "load", lambda self: snapshot)
    fx_calls = []

    def load_fx(self, path):
        fx_calls.append(path)
        return OpexFxTable(2027, {"USD": Decimal("1"), "PEN": Decimal("3.40")}, {"PE": "PEN"})

    monkeypatch.setattr(preparation.OpexFxLoader, "load", load_fx)
    service = preparation.OpexSmartImportPreparationService(tmp_path)
    request = dict(source_path=source, origin="PB", budgeter="TEST", actor="tester")
    pending = service.prepare(**request)

    assert pending.rows == ()
    assert fx_calls == []
    assert pending.decisions[0].cebe_decisions == ()
    assert pending.decisions[0].categoria_gasto == account.categoria_gasto
    review = pending.review_options[0]
    selected = review.cebe_choices[0].options[1]

    ready = service.prepare(**request, overrides=(OpexSmartImportSheetOverride(
        "Licencias", account=review.account_options[0], distribution_mode="IMPORTE",
        cebe_selections=(selected,),
    ),))

    assert len(ready.rows) == 1
    assert ready.rows[0]["sede_cg"] == "SEDE B"
    expected = Decimal(monto).quantize(OPEX_SMART_MONEY_QUANTUM)
    if tipo == "MENSUAL":
        expected *= 12
    assert ready.total_usd == expected
    assert ready.decisions[0].cebe_decisions[0].selected_option == selected
    assert len(fx_calls) == 1

    row = ready.rows[0]
    for group in ("mf", "usd", "ml"):
        monthly = [row[f"{month}_{group}"] for month in MONTHS]
        assert all(value >= 0 for value in monthly)
        assert sum(monthly) == row[f"anio_{group}"]
    workspace = PresupuestoWorkspace(OPEX_MODULE_CONFIG)
    workspace.load(())
    workspace.add_new_rows(ready.rows, description="Importacion inteligente")
    stored = workspace.get_rows()[0]
    for column in OPEX_MODULE_CONFIG.amount_columns:
        assert stored[column] == row[column]
        assert isinstance(stored[column], Decimal)

    batch = build_persistence_batch(workspace, actor="tester", app_version="test")
    payload = json.loads(build_staging_rows(batch)[0].insert_payload)
    for column in OPEX_MODULE_CONFIG.amount_columns:
        assert Decimal(payload[column]) == row[column]

    repository = SimpleNamespace(module_config=OPEX_MODULE_CONFIG, get_all_rows=lambda: (stored,))
    exported = BudgetExcelExportService(repository).export(tmp_path / "precision.xlsx")
    excel = load_workbook(exported.path, data_only=True)
    try:
        sheet = excel[exported.sheet_name]
        values = dict(zip((cell.value for cell in sheet[1]), (cell.value for cell in sheet[2])))
        for column in OPEX_MODULE_CONFIG.amount_columns:
            assert Decimal(str(values[column])).quantize(OPEX_SMART_MONEY_QUANTUM) == row[column]
    finally:
        excel.close()
