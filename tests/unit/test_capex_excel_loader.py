from pathlib import Path

import pytest
from openpyxl import Workbook

from app.config.capex_schema import (
    CAPEX_RAW_TO_INTERNAL,
)
from app.services.capex_excel_loader import (
    CapexWorkbookError,
    build_capex_quality_report,
    load_capex_workbook,
)


pytestmark = pytest.mark.unit


def make_workbook(
    path: Path,
    *,
    year=2026,
    use_usd_alias=False,
    excel_error=False,
    remove_header=None,
):
    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = (
        "PB 2027"
    )

    headers = list(
        CAPEX_RAW_TO_INTERNAL
    )

    if use_usd_alias:
        headers[
            headers.index(
                "3 USD"
            )
        ] = "03 USD"

    if remove_header is not None:
        headers.remove(
            remove_header
        )

    worksheet.cell(
        row=8,
        column=1,
        value=None,
    )

    for index, header in enumerate(
        headers,
        start=2,
    ):
        worksheet.cell(
            row=8,
            column=index,
            value=header,
        )

    row = {
        header: None
        for header in headers
    }

    row.update(
        {
            "Tipo": "PB",
            "A\u00f1o": year,
            "Cantidad": 1,
            "Codigo CEBE":
                "04WF2EAF90",
            "Codigo CECO":
                "04WF2EAF93",
            "01 ML": 100,
            "TOTAL ML": 100,
            "01 USD": 25,
            "TOTAL USD": 25,
        }
    )

    if use_usd_alias:
        row["03 USD"] = 0

    if excel_error:
        row[
            (
                "03 USD"
                if use_usd_alias
                else "3 USD"
            )
        ] = "#N/A"

    for index, header in enumerate(
        headers,
        start=2,
    ):
        worksheet.cell(
            row=9,
            column=index,
            value=row.get(
                header
            ),
        )

    workbook.save(
        path
    )

    workbook.close()


def test_loader_detects_official_structure(
    tmp_path,
):
    path = (
        tmp_path
        / "capex.xlsx"
    )

    make_workbook(
        path
    )

    result = (
        load_capex_workbook(
            path,
            expected_year=None,
        )
    )

    assert result.header_row == 8
    assert result.rows_read == 1
    assert result.valid_count == 1

    clean = (
        result.results[0].row
    )

    assert (
        clean["codigo_cebe"]
        == "04WF2EAF90"
    )

    assert (
        clean["codigo_ceco"]
        == "04WF2EAF93"
    )


def test_loader_skips_blank_rows(
    tmp_path,
):
    path = (
        tmp_path
        / "capex.xlsx"
    )

    make_workbook(
        path
    )

    result = (
        load_capex_workbook(
            path,
            expected_year=None,
        )
    )

    assert result.rows_read == 1


def test_missing_header_is_rejected(
    tmp_path,
):
    path = (
        tmp_path
        / "capex.xlsx"
    )

    make_workbook(
        path,
        remove_header=(
            "Cantidad"
        ),
    )

    with pytest.raises(
        CapexWorkbookError,
        match="Faltantes",
    ):
        load_capex_workbook(
            path,
            expected_year=None,
        )


def test_zero_padded_usd_alias_is_accepted(
    tmp_path,
):
    path = (
        tmp_path
        / "capex.xlsx"
    )

    make_workbook(
        path,
        use_usd_alias=True,
    )

    result = (
        load_capex_workbook(
            path,
            expected_year=None,
        )
    )

    assert result.valid_count == 1


def test_excel_error_marks_row_invalid(
    tmp_path,
):
    path = (
        tmp_path
        / "capex.xlsx"
    )

    make_workbook(
        path,
        excel_error=True,
    )

    result = (
        load_capex_workbook(
            path,
            expected_year=None,
        )
    )

    assert result.invalid_count == 1

    assert any(
        issue.code
        == "INVALID_AMOUNT"
        for issue
        in result.issues
    )


def test_real_year_validation_can_be_enabled(
    tmp_path,
):
    path = (
        tmp_path
        / "capex.xlsx"
    )

    make_workbook(
        path,
        year=2026,
    )

    result = (
        load_capex_workbook(
            path,
            expected_year=2027,
        )
    )

    assert result.invalid_count == 1

    assert any(
        issue.code
        == "YEAR_MISMATCH"
        for issue
        in result.issues
    )


def test_quality_report_contains_summary(
    tmp_path,
):
    path = (
        tmp_path
        / "capex.xlsx"
    )

    make_workbook(
        path
    )

    result = (
        load_capex_workbook(
            path,
            expected_year=None,
        )
    )

    report = (
        build_capex_quality_report(
            result
        )
    )

    assert (
        "Filas leidas: 1"
        in report
    )

    assert (
        "Filas validas: 1"
        in report
    )

    assert (
        "Filas con error: 0"
        in report
    )
