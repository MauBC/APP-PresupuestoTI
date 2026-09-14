from decimal import Decimal

import pytest

from openpyxl import (
    load_workbook,
)

from app.config.budget_modules import (
    CAPEX_MODULE_CONFIG,
    OPEX_MODULE_CONFIG,
)
from app.config.capex_schema import (
    CAPEX_EXCEL_SHEET,
    CAPEX_INTERNAL_TO_RAW,
)
from app.services.budget_excel_export_service import (
    BudgetExcelExportService,
)


pytestmark = pytest.mark.unit


class FakeRepository:
    def __init__(
        self,
        module_config,
        rows,
    ):
        self.module_config = (
            module_config
        )

        self.rows = tuple(
            rows
        )

        self.calls = 0

    def get_all_rows(
        self,
    ):
        self.calls += 1

        return tuple(
            dict(row)
            for row
            in self.rows
        )


def test_opex_exports_business_contract(
    tmp_path,
):
    repository = (
        FakeRepository(
            OPEX_MODULE_CONFIG,
            (
                {
                    "pais":
                        "PER",
                    "ceco":
                        "001234",
                    "presupuestador":
                        "MAURO",
                    "enero_usd":
                        Decimal(
                            "100.25"
                        ),
                    "anio_usd":
                        Decimal(
                            "100.25"
                        ),
                    "row_id":
                        "technical-id",
                    "version":
                        7,
                    "habilitado":
                        True,
                },
            ),
        )
    )

    result = (
        BudgetExcelExportService(
            repository
        )
        .export(
            tmp_path
            / "opex_export"
        )
    )

    assert result.row_count == 1

    assert (
        result.column_count
        ==
        len(
            OPEX_MODULE_CONFIG
            .insert_columns
        )
    )

    assert (
        result.module
        == "OPEX"
    )

    workbook = load_workbook(
        result.path,
        data_only=True,
    )

    try:
        worksheet = (
            workbook[
                "OPEX 2027"
            ]
        )

        headers = tuple(
            cell.value
            for cell
            in worksheet[1]
        )

        assert (
            headers
            ==
            OPEX_MODULE_CONFIG
            .insert_columns
        )

        assert (
            "row_id"
            not in headers
        )

        assert (
            "version"
            not in headers
        )

        assert (
            "habilitado"
            not in headers
        )

        ceco_column = (
            headers.index(
                "ceco"
            )
            + 1
        )

        assert (
            worksheet.cell(
                row=2,
                column=(
                    ceco_column
                ),
            ).value
            == "001234"
        )

    finally:
        workbook.close()


def test_capex_uses_official_excel_headers(
    tmp_path,
):
    repository = (
        FakeRepository(
            CAPEX_MODULE_CONFIG,
            (
                {
                    "tipo":
                        "CAPEX",
                    "pais":
                        "PER",
                    "anio":
                        2027,
                    "presupuestador":
                        "MAURO",
                    "nombre_inversion":
                        "PROYECTO TEST",
                    "codigo_ceco":
                        "04WF2EAF93",
                    "enero_usd":
                        Decimal(
                            "100.123456789"
                        ),
                    "anio_usd":
                        Decimal(
                            "100.123456789"
                        ),
                },
            ),
        )
    )

    result = (
        BudgetExcelExportService(
            repository
        )
        .export(
            tmp_path
            / "capex.xlsx"
        )
    )

    assert (
        result.sheet_name
        == CAPEX_EXCEL_SHEET
    )

    workbook = load_workbook(
        result.path,
        data_only=True,
    )

    try:
        worksheet = (
            workbook[
                CAPEX_EXCEL_SHEET
            ]
        )

        headers = tuple(
            cell.value
            for cell
            in worksheet[1]
        )

        expected = tuple(
            CAPEX_INTERNAL_TO_RAW[
                column
            ]
            for column in (
                CAPEX_MODULE_CONFIG
                .insert_columns
            )
        )

        assert headers == expected

        assert (
            "Presupuestador"
            in headers
        )

        assert (
            "Nombre de Inversión"
            in headers
        )

    finally:
        workbook.close()


def test_export_preserves_string_codes(
    tmp_path,
):
    repository = (
        FakeRepository(
            CAPEX_MODULE_CONFIG,
            (
                {
                    "codigo_ceco":
                        "001234",
                    "codigo_cebe":
                        "000456",
                },
            ),
        )
    )

    result = (
        BudgetExcelExportService(
            repository
        )
        .export(
            tmp_path
            / "codes.xlsx"
        )
    )

    workbook = load_workbook(
        result.path,
        data_only=True,
    )

    try:
        worksheet = (
            workbook[
                CAPEX_EXCEL_SHEET
            ]
        )

        headers = tuple(
            cell.value
            for cell
            in worksheet[1]
        )

        ceco = (
            headers.index(
                CAPEX_INTERNAL_TO_RAW[
                    "codigo_ceco"
                ]
            )
            + 1
        )

        cebe = (
            headers.index(
                CAPEX_INTERNAL_TO_RAW[
                    "codigo_cebe"
                ]
            )
            + 1
        )

        assert (
            worksheet.cell(
                2,
                ceco,
            ).value
            == "001234"
        )

        assert (
            worksheet.cell(
                2,
                cebe,
            ).value
            == "000456"
        )

    finally:
        workbook.close()


def test_empty_database_still_exports_headers(
    tmp_path,
):
    repository = (
        FakeRepository(
            OPEX_MODULE_CONFIG,
            (),
        )
    )

    result = (
        BudgetExcelExportService(
            repository
        )
        .export(
            tmp_path
            / "empty"
        )
    )

    assert result.row_count == 0

    workbook = load_workbook(
        result.path
    )

    try:
        worksheet = (
            workbook[
                "OPEX 2027"
            ]
        )

        assert (
            worksheet.max_row
            == 1
        )

    finally:
        workbook.close()


def test_export_reads_repository_once(
    tmp_path,
):
    repository = (
        FakeRepository(
            OPEX_MODULE_CONFIG,
            (),
        )
    )

    BudgetExcelExportService(
        repository
    ).export(
        tmp_path
        / "once.xlsx"
    )

    assert repository.calls == 1


def test_disabled_rows_are_not_exported(
    tmp_path,
):
    repository = (
        FakeRepository(
            OPEX_MODULE_CONFIG,
            (
                {
                    "pais": "PER",
                    "habilitado": True,
                },
                {
                    "pais": "CHL",
                    "habilitado": False,
                },
            ),
        )
    )

    result = (
        BudgetExcelExportService(
            repository
        )
        .export(
            tmp_path
            / "active_only.xlsx"
        )
    )

    assert result.row_count == 1

    workbook = load_workbook(
        result.path,
        data_only=True,
    )

    try:
        worksheet = workbook[
            "OPEX 2027"
        ]

        headers = tuple(
            cell.value
            for cell
            in worksheet[1]
        )

        pais_column = (
            headers.index(
                "pais"
            )
            + 1
        )

        assert (
            worksheet.cell(
                row=2,
                column=pais_column,
            ).value
            == "PER"
        )

        assert worksheet.max_row == 2

    finally:
        workbook.close()
