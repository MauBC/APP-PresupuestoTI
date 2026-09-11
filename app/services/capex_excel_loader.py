from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Any
import unicodedata

from openpyxl import load_workbook

from app.config.capex_schema import (
    CAPEX_EXCEL_SHEET,
    CAPEX_HEADER_ALIASES,
    CAPEX_RAW_TO_INTERNAL,
)
from app.models.capex_import import (
    CapexCleanRowResult,
    CapexIssueSeverity,
)
from app.services.capex_validator import (
    clean_and_validate_capex_row,
)


class CapexWorkbookError(
    ValueError
):
    pass


@dataclass(
    frozen=True,
    slots=True,
)
class CapexExcelLoadResult:
    source_path: str
    sheet_name: str
    header_row: int
    rows_read: int
    results: tuple[
        CapexCleanRowResult,
        ...
    ]

    @property
    def valid_results(
        self,
    ) -> tuple[
        CapexCleanRowResult,
        ...
    ]:
        return tuple(
            result
            for result in self.results
            if result.is_valid
        )

    @property
    def invalid_results(
        self,
    ) -> tuple[
        CapexCleanRowResult,
        ...
    ]:
        return tuple(
            result
            for result in self.results
            if not result.is_valid
        )

    @property
    def valid_count(
        self,
    ) -> int:
        return len(
            self.valid_results
        )

    @property
    def invalid_count(
        self,
    ) -> int:
        return len(
            self.invalid_results
        )

    @property
    def issues(
        self,
    ):
        return tuple(
            issue
            for result in self.results
            for issue in result.issues
        )

    @property
    def error_count(
        self,
    ) -> int:
        return sum(
            1
            for issue in self.issues
            if (
                issue.severity
                == CapexIssueSeverity.ERROR
            )
        )

    @property
    def warning_count(
        self,
    ) -> int:
        return sum(
            1
            for issue in self.issues
            if (
                issue.severity
                == CapexIssueSeverity.WARNING
            )
        )

    @property
    def info_count(
        self,
    ) -> int:
        return sum(
            1
            for issue in self.issues
            if (
                issue.severity
                == CapexIssueSeverity.INFO
            )
        )


def normalize_capex_header(
    value,
) -> str:
    if value is None:
        return ""

    text = unicodedata.normalize(
        "NFC",
        str(value),
    )

    text = " ".join(
        text.strip().split()
    )

    return (
        CAPEX_HEADER_ALIASES
        .get(
            text,
            text,
        )
    )


def detect_capex_header_row(
    rows,
    *,
    scan_rows: int = 20,
) -> int:
    expected = set(
        CAPEX_RAW_TO_INTERNAL
    )

    best_row = None
    best_score = -1

    upper_row = min(
        len(rows),
        scan_rows,
    )

    for row_index in range(
        upper_row
    ):
        headers = {
            normalize_capex_header(
                value
            )
            for value
            in rows[row_index]
        }

        headers.discard(
            ""
        )

        score = len(
            headers
            & expected
        )

        if score > best_score:
            best_score = score
            best_row = (
                row_index + 1
            )

    if (
        best_row is None
        or best_score < 40
    ):
        raise CapexWorkbookError(
            "No se pudo identificar "
            "la fila de encabezados CAPEX."
        )

    return best_row


def read_capex_headers(
    rows,
    *,
    header_row: int,
) -> tuple[
    str,
    ...
]:
    return tuple(
        normalize_capex_header(
            value
        )
        for value
        in rows[
            header_row - 1
        ]
    )


def validate_capex_headers(
    headers,
):
    non_empty = tuple(
        header
        for header in headers
        if header
    )

    counts = Counter(
        non_empty
    )

    duplicates = tuple(
        header
        for header, count
        in counts.items()
        if count > 1
    )

    if duplicates:
        raise CapexWorkbookError(
            "Encabezados CAPEX duplicados: "
            + ", ".join(
                sorted(
                    duplicates
                )
            )
        )

    expected = set(
        CAPEX_RAW_TO_INTERNAL
    )

    actual = set(
        non_empty
    )

    missing = sorted(
        expected - actual
    )

    unexpected = sorted(
        actual - expected
    )

    if missing or unexpected:
        details = []

        if missing:
            details.append(
                "Faltantes: "
                + ", ".join(
                    missing
                )
            )

        if unexpected:
            details.append(
                "No reconocidas: "
                + ", ".join(
                    unexpected
                )
            )

        raise CapexWorkbookError(
            "La estructura de PB 2027 "
            "no coincide con el contrato. "
            + " | ".join(
                details
            )
        )

    if len(actual) != 50:
        raise CapexWorkbookError(
            "Se esperaban exactamente "
            "50 columnas CAPEX de negocio."
        )


def _row_is_empty(
    values: dict[str, Any],
) -> bool:
    for value in values.values():
        if value is None:
            continue

        if (
            isinstance(
                value,
                str,
            )
            and not value.strip()
        ):
            continue

        return False

    return True


def load_capex_workbook(
    source_path,
    *,
    expected_year: int | None,
    sheet_name: str = CAPEX_EXCEL_SHEET,
) -> CapexExcelLoadResult:
    source = Path(
        source_path
    )

    if not source.exists():
        raise CapexWorkbookError(
            f"No existe el archivo: {source}"
        )

    workbook = load_workbook(
        source,
        read_only=True,
        data_only=True,
    )

    try:
        if (
            sheet_name
            not in workbook.sheetnames
        ):
            raise CapexWorkbookError(
                "No existe la hoja "
                f"{sheet_name!r}."
            )

        worksheet = workbook[
            sheet_name
        ]

        sheet_rows = list(
            worksheet.iter_rows(
                values_only=True
            )
        )

        header_row = (
            detect_capex_header_row(
                sheet_rows
            )
        )

        headers = read_capex_headers(
            sheet_rows,
            header_row=header_row,
        )

        validate_capex_headers(
            headers
        )

        results = []

        for (
            row_number,
            values,
        ) in enumerate(
            sheet_rows[
                header_row:
            ],
            start=(
                header_row + 1
            ),
        ):
            raw_row = {
                header: value
                for header, value
                in zip(
                    headers,
                    values,
                )
                if header
            }

            if _row_is_empty(
                raw_row
            ):
                continue

            result = (
                clean_and_validate_capex_row(
                    raw_row,
                    row_number=(
                        row_number
                    ),
                    expected_year=(
                        expected_year
                    ),
                )
            )

            results.append(
                result
            )

        return CapexExcelLoadResult(
            source_path=str(
                source.resolve()
            ),
            sheet_name=sheet_name,
            header_row=header_row,
            rows_read=len(
                results
            ),
            results=tuple(
                results
            ),
        )

    finally:
        workbook.close()


def build_capex_quality_report(
    result: CapexExcelLoadResult,
) -> str:
    issue_code_counts = Counter(
        issue.code
        for issue in result.issues
    )

    issue_column_counts = Counter(
        issue.column
        for issue in result.issues
    )

    lines = [
        "=" * 100,
        "M6 - CALIDAD DE IMPORTACION CAPEX",
        "=" * 100,
        "",
        f"Archivo: {result.source_path}",
        f"Hoja: {result.sheet_name}",
        f"Fila encabezados: {result.header_row}",
        "",
        "RESUMEN",
        "-" * 100,
        f"Filas leidas: {result.rows_read:,}",
        f"Filas validas: {result.valid_count:,}",
        f"Filas con error: {result.invalid_count:,}",
        f"Errores: {result.error_count:,}",
        f"Advertencias: {result.warning_count:,}",
        f"Informativos: {result.info_count:,}",
        "",
        "ISSUES POR CODIGO",
        "-" * 100,
    ]

    if issue_code_counts:
        for code, count in (
            issue_code_counts
            .most_common()
        ):
            lines.append(
                f"{code}: {count:,}"
            )
    else:
        lines.append(
            "[OK] No se detectaron issues."
        )

    lines.extend(
        [
            "",
            "ISSUES POR COLUMNA",
            "-" * 100,
        ]
    )

    if issue_column_counts:
        for column, count in (
            issue_column_counts
            .most_common()
        ):
            lines.append(
                f"{column}: {count:,}"
            )
    else:
        lines.append(
            "[OK] No se detectaron issues."
        )

    lines.extend(
        [
            "",
            "DETALLE",
            "-" * 100,
        ]
    )

    if not result.issues:
        lines.append(
            "[OK] Todas las filas "
            "pasaron sin issues."
        )

    else:
        for issue in result.issues:
            lines.append(
                f"Fila {issue.row_number} | "
                f"{issue.severity.value} | "
                f"{issue.code} | "
                f"{issue.column} | "
                f"{issue.message} | "
                f"raw={issue.raw_value!r}"
            )

    return (
        "\n".join(
            lines
        )
        + "\n"
    )
