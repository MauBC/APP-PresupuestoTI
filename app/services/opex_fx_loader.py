from decimal import (
    Decimal,
    InvalidOperation,
    ROUND_HALF_UP,
)
from pathlib import Path
import unicodedata

from openpyxl import load_workbook

from app.models.opex_fx import (
    OpexFxTable,
)


TC_QUANTUM = Decimal("0.01")


class OpexFxLoadError(
    ValueError
):
    pass


COUNTRY_CODES = {
    "PERU": "PE",
    "BOLIVIA": "BO",
    "ECUADOR": "EC",
    "COLOMBIA": "CO",
    "EL SALVADOR": "SV",
    "HONDURAS": "HN",
    "GUATEMALA": "GT",
    "PANAMA": "PA",
    "NICARAGUA": "NI",
    "MEXICO": "MX",
    "COSTA RICA": "CR",
    "CHILE": "CL",
}


def _normalize_text(
    value,
):
    text = str(
        value
        if value is not None
        else ""
    ).strip()

    text = "".join(
        char
        for char in unicodedata.normalize(
            "NFD",
            text,
        )
        if unicodedata.category(
            char
        ) != "Mn"
    )

    return " ".join(
        text.upper().split()
    )


class OpexFxLoader:
    def load(
        self,
        path,
        *,
        year=2027,
    ) -> OpexFxTable:
        source = Path(
            path
        )

        if not source.exists():
            raise OpexFxLoadError(
                f"No existe el archivo TC: {source}"
            )

        workbook = load_workbook(
            source,
            read_only=True,
            data_only=True,
        )

        matches = []

        for sheet in workbook.worksheets:
            for row_number, row in enumerate(
                sheet.iter_rows(
                    values_only=True
                ),
                start=1,
            ):
                normalized = tuple(
                    _normalize_text(
                        value
                    )
                    for value in row
                )

                if (
                    "PAIS" in normalized
                    and "MONEDA" in normalized
                    and "TC" in normalized
                ):
                    matches.append(
                        (
                            sheet,
                            row_number,
                            normalized,
                        )
                    )

        if len(matches) != 1:
            raise OpexFxLoadError(
                "TC.xlsx debe contener una sola "
                "tabla con encabezados "
                "Pais, Moneda y TC."
            )

        (
            sheet,
            header_row,
            headers,
        ) = matches[0]

        indexes = {
            name:
                headers.index(
                    name
                )
            for name in (
                "PAIS",
                "MONEDA",
                "TC",
            )
        }

        rates = {
            "USD":
                Decimal("1.00"),
        }

        local_currency = {}

        for excel_row, values in enumerate(
            sheet.iter_rows(
                min_row=header_row + 1,
                values_only=True,
            ),
            start=header_row + 1,
        ):
            pais_raw = (
                values[
                    indexes["PAIS"]
                ]
            )

            moneda_raw = (
                values[
                    indexes["MONEDA"]
                ]
            )

            tc_raw = (
                values[
                    indexes["TC"]
                ]
            )

            if (
                pais_raw is None
                and moneda_raw is None
                and tc_raw is None
            ):
                continue

            pais_name = (
                _normalize_text(
                    pais_raw
                )
            )

            moneda = (
                _normalize_text(
                    moneda_raw
                )
            )

            if not pais_name:
                raise OpexFxLoadError(
                    f"Fila {excel_row}: Pais vacio."
                )

            if pais_name not in COUNTRY_CODES:
                raise OpexFxLoadError(
                    f"Fila {excel_row}: pais no "
                    f"reconocido: {pais_raw}."
                )

            country = (
                COUNTRY_CODES[
                    pais_name
                ]
            )

            if not moneda:
                raise OpexFxLoadError(
                    f"Fila {excel_row}: Moneda vacia."
                )

            try:
                tc = Decimal(
                    str(tc_raw)
                ).quantize(
                    TC_QUANTUM,
                    rounding=ROUND_HALF_UP,
                )

            except (
                InvalidOperation,
                TypeError,
                ValueError,
            ) as exc:
                raise OpexFxLoadError(
                    f"Fila {excel_row}: TC invalido."
                ) from exc

            if tc <= Decimal("0"):
                raise OpexFxLoadError(
                    f"Fila {excel_row}: TC debe "
                    "ser mayor que cero."
                )

            if (
                moneda == "USD"
                and tc != Decimal("1.00")
            ):
                raise OpexFxLoadError(
                    f"Fila {excel_row}: {pais_raw} "
                    "declara moneda USD pero su "
                    f"TC es {tc}; USD debe usar "
                    "TC 1.00."
                )

            if country in local_currency:
                raise OpexFxLoadError(
                    f"Pais duplicado en TC.xlsx: "
                    f"{pais_raw}."
                )

            current_rate = rates.get(
                moneda
            )

            if (
                current_rate is not None
                and current_rate != tc
            ):
                raise OpexFxLoadError(
                    f"La moneda {moneda} tiene "
                    "mas de un TC distinto."
                )

            rates[
                moneda
            ] = tc

            local_currency[
                country
            ] = moneda

        if not local_currency:
            raise OpexFxLoadError(
                "TC.xlsx no contiene registros."
            )

        return OpexFxTable(
            year=year,
            currency_per_usd=rates,
            local_currency_by_country=(
                local_currency
            ),
        )
